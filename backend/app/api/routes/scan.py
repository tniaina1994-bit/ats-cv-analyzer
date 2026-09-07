from fastapi import APIRouter, UploadFile, File, Form
from fastapi.responses import JSONResponse
from typing import List, Set
import io
import re
import time

router = APIRouter()

TECH_SKILLS = [
    "python", "java", "javascript", "typescript", "angular", "react", "vue",
    "node.js", "php", "c++", "c#", "ruby", "go", "golang", "rust", "swift",
    "kotlin", "scala", "dart", "elixir", "haskell", "perl", "lua", "r",
    "docker", "kubernetes", "k8s", "ansible", "terraform", "jenkins",
    "gitlab", "github", "webpack", "vite", "npm", "yarn", "pip",
    "html", "css", "scss", "sass", "less", "tailwind", "bootstrap",
    "mysql", "postgresql", "postgres", "mongodb", "redis", "elasticsearch",
    "sqlite", "mariadb", "oracle", "cassandra", "dynamodb", "neo4j",
    "aws", "azure", "gcp", "heroku", "digitalocean", "cloudflare", "vercel", "netlify",
    "linux", "ubuntu", "debian", "centos", "redhat", "fedora", "windows",
    "rest", "graphql", "grpc", "api", "soap", "websocket",
    "agile", "scrum", "kanban", "sprint", "jira", "confluence", "trello",
    "ci/cd", "devops", "sre", "pipeline",
    "tensorflow", "pytorch", "keras", "pandas", "numpy", "opencv",
    "spacy", "nltk", "scikit-learn",
    "django", "flask", "fastapi", "spring", "laravel", "rails", "express",
    "next.js", "nuxt", "nestjs",
    "figma", "sketch", "photoshop", "illustrator",
    "flutter", "react native", "xamarin", "maui", "ionic",
    "seo", "google analytics", "google ads",
    "jest", "mocha", "pytest", "selenium", "cypress", "playwright",
    "git", "svn",
    "active directory", "oauth", "jwt", "sso", "saml", "ldap",
    "vpn", "firewall", "dns", "ssh", "ssl", "tls",
    "proxmox", "vmware", "hyper-v", "kvm",
    "nginx", "apache", "tomcat", "iis",
    "kafka", "rabbitmq", "celery", "nats",
    "excel", "power bi", "tableau",
    "shopify", "woocommerce", "magento", "wordpress",
    "cloudflare", "cdn",
    "notion", "slack", "teams", "discord",
    "openai", "langchain", "llm", "nlp", "rag",
    "machine learning", "deep learning", "ai",
    "sql", "nosql", "plsql",
    "svg", "xml", "json", "yaml",
    "prometheus", "grafana", "datadog", "sentry", "zabbix",
    "hadoop", "spark", "flink", "airflow",
    "databricks", "snowflake", "bigquery", "redshift",
    "jupyter", "colab",
    "unity", "unreal", "godot", "blender",
    "web3", "blockchain", "solidity",
    "traefik", "haproxy", "envoy", "istio",
    "podman", "containerd",
    "pulumi", "cloudformation", "helm", "kustomize",
    "visual studio", "vscode", "intellij", "pycharm",
    "android studio", "xcode",
    "maya", "cinema 4d", "premiere", "after effects", "davinci",
    "salesforce", "hubspot", "zendesk",
    "stripe", "paypal",
    "strapi", "payload", "directus",
    "openapi", "swagger", "postman",
    "pwa", "webassembly", "wasm",
    "three.js", "d3.js", "chart.js",
    "microservices", "monolith", "serverless",
    "cqrs", "event sourcing", "event driven",
    "solid", "clean code", "refactoring",
    "pair programming", "code review",
    "load testing", "stress testing", "benchmark",
    "lighthouse", "core web vitals",
    "owasp", "gdpr", "hipaa", "soc 2", "iso 27001",
    "product management", "product owner",
    "scrum master", "agile coach",
    "tech lead", "engineering manager", "cto",
    "architect", "solution architect",
    "data engineer", "data scientist",
    "ml engineer", "ai engineer",
    "backend developer", "frontend developer", "full stack",
    "mobile developer", "cloud engineer", "security engineer",
    "qa engineer", "test engineer",
    "ux designer", "ui designer", "ux researcher",
    "technical writer", "devrel",
    "firestore", "supabase", "appwrite", "firebase",
    "etcd", "zookeeper", "consul", "vault",
    "packer", "vagrant",
    "openshift", "rancher", "linkerd",
    "minikube", "k3s", "microk8s",
    "gunicorn", "uvicorn",
    "pydantic", "sqlalchemy", "alembic",
    "keycloak", "auth0", "okta",
    "sonarqube", "eslint", "prettier", "mypy",
    "tox", "coverage",
    "circleci", "travis", "drone", "buildkite",
    "bash", "zsh", "powershell",
    "vb.net", "asp.net",
    "couchdb", "couchbase", "scylla", "cosmosdb",
    "memcached",
    "ghost", "bigcommerce", "prestashop",
    "sinatra", "starlette", "fastify", "koa",
    "symfony", "codeigniter",
    "springboot", "quarkus", "micronaut",
    "linode", "vultr",
    "lambda", "rds", "sqs", "sns", "ec2", "s3", "cloudwatch",
    "gke", "aks", "eks",
    "logstash", "filebeat", "fluentd",
    "newrelic", "bugsnag", "rollbar",
    "pagerduty", "opsgenie",
    "d3", "leaflet", "mapbox",
    "mapbox", "webrtc",
    "hls", "dash",
    "ffmpeg", "gstreamer",
    "podcast", "streaming",
    "opentelemetry", "otel",
    "certbot",
    "docker compose",
    "junior", "senior", "mid", "staff", "principal",
    "intern", "trainee",
    "manager", "director", "lead",
    "consultant", "analyst", "specialist",
    "mongodb", "sql server",
    "bitbucket",
    "terraform", "ansible", "puppet", "chef",
    "github actions", "gitlab ci",
    "argocd", "tekton", "flux",
    "spinnaker",
    "story point", "velocity", "backlog",
    "retrospective", "standup",
    "burndown", "burnup",
    "lean", "six sigma",
    "itil", "cobit",
    "togaf",
    "wireframe", "prototype", "usability",
    "accessibility", "wcag",
    "responsive", "mobile first",
    "spa", "ssr", "ssg", "isr",
    "monorepo",
    "feature flag", "a/b testing",
    "funnel", "retention", "churn",
    "kpi", "metric", "dashboard",
    "conversion", "revenue",
    "nps", "csat",
    "okr",
    "portfolio", "linkedin",
    "mentorship", "coaching",
    "leadership", "communication",
    "remote work", "freelancing",
    "startup",
    "etl", "elt",
    "data pipeline", "data lake", "data warehouse",
    "apache kafka", "apache spark", "apache airflow",
    "dbt",
    "mlops", "dataops",
    "transformer", "bert", "gpt",
    "stable diffusion", "midjourney",
    "ollama", "vllm",
    "langchain", "llamaindex",
    "maven", "gradle",
    "power shell",
    "migration de données",
    "tcp/ip",
]

TECH_SET = set(s.lower() for s in TECH_SKILLS)


def extract_skills_from_job(job_content: str) -> List[str]:
    """Extract skills from job offer text"""
    found: Set[str] = set()
    content_lower = job_content.lower()

    for skill in TECH_SKILLS:
        pattern = r'\b' + re.escape(skill).replace(r'\ ', r'\s*') + r'\b'
        if re.search(pattern, content_lower):
            found.add(skill.upper())

    return list(found)


def match_skill_in_cv(skill: str, cv_text_lower: str) -> bool:
    """Check if skill appears in CV"""
    pattern = r'\b' + re.escape(skill.lower()).replace(r'\ ', r'\s*') + r'\b'
    return bool(re.search(pattern, cv_text_lower))


def extract_personal_info(cv_text: str) -> dict:
    """Extract personal information from CV text"""
    cv_lower = cv_text.lower()
    first_lines = cv_text[:500]

    # Email
    email_match = re.search(r'[\w\.\-+]+@[\w\.\-]+\.\w+', cv_text)
    email = email_match.group(0) if email_match else ""

    # Phone
    phone_match = re.search(r'[\+]?[\d\s\-\(\)]{8,}', cv_text)
    phone = phone_match.group(0).strip() if phone_match else ""

    # LinkedIn
    linkedin_match = re.search(r'linkedin\.com/in/[\w\-\%\._]+', cv_text, re.IGNORECASE)
    linkedin = linkedin_match.group(0) if linkedin_match else ""

    # GitHub
    github_match = re.search(r'github\.com/[\w\-]+', cv_text, re.IGNORECASE)
    github = github_match.group(0) if github_match else ""

    # Website/portfolio
    website_match = re.search(r'(?:https?://)?(?!linkedin|github)[\w\-]+\.(?:com|fr|io|dev|net|org)(?:/[\w\-\./]*)?', cv_text, re.IGNORECASE)
    website = ""
    if website_match:
        w = website_match.group(0)
        if "linkedin" not in w.lower() and "github" not in w.lower():
            website = w

    # Location / City
    location = ""
    city_patterns = [
        r'(?:located|situé|située|basé|basée|ville|city)\s*[:]\s*([A-ZÀ-Ÿa-zÀ-ÿ\s\-]+)',
        r'([A-ZÀ-Ÿ][a-zÀ-ÿ]+(?:-[A-ZÀ-Ÿ][a-zÀ-ÿ]+)?)\s*,?\s*(?:France|Madagascar|Tunisie|Maroc|Cameroun|Sénégal|Belgique|Canada|Allemagne)',
    ]
    for p in city_patterns:
        m = re.search(p, cv_text)
        if m:
            location = m.group(0).strip()
            break

    # Languages
    lang_keywords = ["français", "french", "english", "anglais", "malagasy", "malgache",
                     "espagnol", "spanish", "allemand", "german", "arabe", "arabic",
                     "chinois", "chinese", "japonais", "japanese", "portugais", "portuguese"]
    languages = [k.title() for k in lang_keywords if k in cv_lower]

    # Driving license
    has_driving_license = bool(re.search(r'(permis\s*[ab]|driving\s*license|licence\s+de\s+conduire)', cv_lower))

    # Nationality
    nationality = ""
    nat_match = re.search(r'(?:nationalité|nationality|nationalité)\s*[:]\s*([A-ZÀ-Ÿa-zÀ-ÿ\s]+)', cv_text, re.IGNORECASE)
    if nat_match:
        nationality = nat_match.group(1).strip()

    # Name heuristic: first non-empty line that is short and doesn't contain @ or digits
    name = ""
    for line in cv_text.split('\n'):
        line = line.strip()
        if (2 < len(line) < 50
            and '@' not in line
            and not any(c.isdigit() for c in line)
            and not re.search(r'(cv|resume|curriculum|expérience|compétence|formation|skills|email|téléphone|phone|adresse)', line, re.IGNORECASE)):
            name = line
            break

    return {
        "name": name,
        "email": email,
        "phone": phone,
        "linkedin": linkedin,
        "github": github,
        "website": website,
        "location": location,
        "languages": languages,
        "has_driving_license": has_driving_license,
        "nationality": nationality,
    }


def analyze_cv_content(cv_text: str, job_skills: List[str]) -> dict:
    """Analyze CV against job skills"""
    cv_lower = cv_text.lower()

    # Detect ALL skills present in CV
    cv_skills_found = []
    for skill in TECH_SKILLS:
        pattern = r'\b' + re.escape(skill).replace(r'\ ', r'\s*') + r'\b'
        if re.search(pattern, cv_lower):
            cv_skills_found.append(skill.upper())

    matched = [s for s in job_skills if match_skill_in_cv(s, cv_lower)]
    missing = [s for s in job_skills if not match_skill_in_cv(s, cv_lower)]

    skill_score = round((len(matched) / len(job_skills)) * 100) if job_skills else 0

    # Experience detection
    exp_patterns = [
        r'(\d+)\s*(?:ans|years?|années?)',
        r'expérience\s*:\s*(\d+)',
        r'expérience\s+de\s+(\d+)',
    ]
    years = 0
    for p in exp_patterns:
        m = re.search(p, cv_lower)
        if m:
            years = max(years, int(m.group(1)))

    exp_score = 100 if years >= 5 else 75 if years >= 3 else 50 if years >= 1 else 30

    # Education detection
    edu_keywords = ["master", "bac+5", "bac+4", "licence", "bachelor", "diplome", "degree", "ingénieur", "université", "university"]
    edu_found = [k for k in edu_keywords if k in cv_lower]
    edu_score = 80 if edu_found else 0

    # ATS checks
    ats_checks = {
        "text_extractable": bool(cv_text.strip()),
        "has_email": bool(re.search(r'[\w\.-]+@[\w\.-]+\.\w+', cv_text)),
        "has_phone": bool(re.search(r'[\+]?[\d\s\-\(\)]{8,}', cv_text)),
        "has_sections": bool(re.search(r'(exp.rience|formation|comp.tence|skills|experience)', cv_lower)),
    }
    ats_score = round(sum(ats_checks.values()) / len(ats_checks) * 100)

    global_score = round(skill_score * 0.5 + exp_score * 0.2 + edu_score * 0.1 + ats_score * 0.2)

    recommendations = []
    if not ats_checks["has_email"]:
        recommendations.append("Ajoutez votre email")
    if not ats_checks["has_phone"]:
        recommendations.append("Ajoutez votre telephone")
    if not ats_checks["has_sections"]:
        recommendations.append("Structurez votre CV avec des sections")
    if missing:
        recommendations.append(f"Competences manquantes : {', '.join(missing[:3])}")

    personal_info = extract_personal_info(cv_text)

    return {
        "global_score": global_score,
        "scores": {"skills": skill_score, "experience": exp_score, "education": edu_score, "ats_quality": ats_score},
        "personal_info": personal_info,
        "cv_skills_detected": cv_skills_found,
        "matched_skills": matched,
        "missing_skills": missing,
        "ats_checks": ats_checks,
        "experience_years": years,
        "education_found": edu_found,
        "recommendations": recommendations,
    }


@router.post("/scan")
async def scan_cv(
    file: UploadFile = File(...),
    job_offer: str = Form(...),
):
    t0 = time.time()

    allowed_types = [".pdf", ".docx"]
    filename = file.filename or ""
    file_ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""

    if file_ext not in allowed_types:
        return JSONResponse(status_code=400, content={"error": f"Formats acceptes : {allowed_types}"})

    content = await file.read()
    if len(content) > 10 * 1024 * 1024:
        return JSONResponse(status_code=400, content={"error": "Fichier trop volumineux (max 10MB)"})

    cv_text = ""
    ocr_used = False

    if file_ext == ".pdf":
        try:
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(content))
            for page in reader.pages:
                text = page.extract_text()
                if text:
                    cv_text += text + "\n"
        except Exception as e:
            return JSONResponse(status_code=400, content={"error": f"Erreur PDF : {str(e)}"})

        if not cv_text.strip():
            ocr_text = ocr_from_pdf(content)
            if ocr_text.strip():
                cv_text = ocr_text
                ocr_used = True

    elif file_ext == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(content))
            for para in doc.paragraphs:
                if para.text.strip():
                    cv_text += para.text + "\n"
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            cv_text += cell.text + "\n"
        except Exception as e:
            return JSONResponse(status_code=400, content={"error": f"Erreur DOCX : {str(e)}"})

    if not cv_text.strip():
        return JSONResponse(status_code=400, content={"error": "Impossible d'extraire le texte du CV."})

    job_skills = extract_skills_from_job(job_offer)
    result = analyze_cv_content(cv_text, job_skills)

    elapsed = round(time.time() - t0, 3)

    return {
        "filename": filename,
        "ocr_used": ocr_used,
        "cv_text_preview": cv_text[:500] + "..." if len(cv_text) > 500 else cv_text,
        "job_skills_found": job_skills,
        "analysis": result,
        "processing_time_ms": elapsed * 1000,
    }


def ocr_from_pdf(content: bytes) -> str:
    try:
        from pypdf import PdfReader
        from PIL import Image
        import pytesseract
        reader = PdfReader(io.BytesIO(content))
        full_text = ""
        for page in reader.pages:
            if "/XObject" in (page.get("/Resources") or {}):
                x_objects = page["/Resources"]["/XObject"].get_object()
                for obj_name in x_objects:
                    x_obj = x_objects[obj_name].get_object()
                    if x_obj.get("/Subtype") == "/Image":
                        w, h = x_obj["/Width"], x_obj["/Height"]
                        data = x_obj.get_data()
                        try:
                            img = Image.frombytes("RGB", (w, h), data)
                            full_text += pytesseract.image_to_string(img, lang="fra+eng") + "\n"
                        except Exception:
                            try:
                                img = Image.frombytes("L", (w, h), data)
                                full_text += pytesseract.image_to_string(img, lang="fra+eng") + "\n"
                            except Exception:
                                continue
        return full_text
    except Exception:
        return ""